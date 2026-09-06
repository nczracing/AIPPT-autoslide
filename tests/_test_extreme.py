# -*- coding: utf-8 -*-
"""极端边界压测：非常规比例图片 / 极长标题 / 空内容页

装饰形状（无文字非图片）允许出血式越界（有意裁切设计），
边界与贴边检查只约束含文字形状和图片。
"""
import os
import sys
import tempfile

sys.path.insert(0, '.')

from PIL import Image as PILImage, ImageDraw
from settings_module import Presentation, Slide
from core.pptx_builder import PPTXBuilder
from pptx import Presentation as PptxPresentation

SW, SH = 13.333, 7.5


def make_img(directory, name, w, h):
    p = os.path.join(directory, name)
    img = PILImage.new('RGB', (w, h), (90, 120, 160))
    d = ImageDraw.Draw(img)
    for i in range(0, w, 60):
        d.rectangle([i, 0, i + 30, h], fill=(120, 150, 190))
    img.save(p)
    return p


def main():
    tmp = tempfile.mkdtemp()
    extreme = {
        'pano': make_img(tmp, 'pano.png', 2000, 400),        # 5:1 全景
        'ultratall': make_img(tmp, 'ultratall.png', 300, 1200),  # 1:4 超竖
        'tiny': make_img(tmp, 'tiny.png', 50, 50),           # 极小图
        'huge': make_img(tmp, 'huge.png', 3000, 2400),       # 超大图 5:4
    }

    p = Presentation(title='极端边界测试', theme='business', language='zh')
    cases = [
        Slide(page=1, title='全景图right', points=['全景比例测试要点一', '验证5:1宽图在右侧的缩放'],
              detail='全景图极端比例验证。', layout='title_content',
              image_path=extreme['pano'], image_layout='right'),
        Slide(page=2, title='全景图top', points=['全景比例测试要点一'],
              detail='全景图顶部布局验证。', layout='title_content',
              image_path=extreme['pano'], image_layout='top'),
        Slide(page=3, title='超竖图left', points=['超竖比例测试要点一', '验证1:4高图在左侧'],
              detail='超竖图极端比例验证。', layout='title_content',
              image_path=extreme['ultratall'], image_layout='left'),
        Slide(page=4, title='极小图right', points=['极小图片测试要点一'],
              detail='极小图片不应被放大到失真或占据异常空间。', layout='title_content',
              image_path=extreme['tiny'], image_layout='right'),
        Slide(page=5, title='超大图right', points=['超大图片测试要点一'],
              detail='超大图片按比例约束缩放。', layout='title_content',
              image_path=extreme['huge'], image_layout='right'),
        Slide(page=6, title='这是一个非常非常非常长的标题用于测试标题自适应字号机制是否能够正确处理极端长度的标题文本而不会越出幻灯片边界造成文字贴边或者被截断的问题出现',
              points=['长标题页面要点一', '长标题页面要点二'], detail='长标题压缩验证。',
              layout='title_content'),
        Slide(page=7, title='空内容页', points=[], detail='', layout='title_content'),
    ]
    for s in cases:
        p.add_slide(s)
    out = os.path.join(tmp, 'extreme.pptx')
    PPTXBuilder().build(p, out)

    prs = PptxPresentation(out)
    issues = []
    for idx, s in enumerate(prs.slides, 1):
        pics, texts = [], []
        for sh in s.shapes:
            try:
                l, t, w, h = sh.left.inches, sh.top.inches, sh.width.inches, sh.height.inches
            except Exception:
                continue
            has_text = sh.has_text_frame and sh.text_frame.text.strip()
            is_pic = sh.shape_type == 13 or sh.__class__.__name__ == 'Picture'
            if (has_text or is_pic) and (
                l < -0.05 or t < -0.05 or l + w > SW + 0.05 or t + h > SH + 0.05
            ):
                issues.append(f'页{idx} 越界 {sh.shape_type} ({l:.2f},{t:.2f},{w:.2f},{h:.2f})')
            if has_text:
                m = (l, t, SW - l - w, SH - t - h)
                if min(m) < 0.25:
                    issues.append(
                        f'页{idx} 文字贴边 {sh.text_frame.text[:12]!r} '
                        f'边距{[round(x, 2) for x in m]}'
                    )
                if not (sh.is_placeholder and sh.placeholder_format.idx == 0):
                    texts.append(((l, t, w, h), sh.text_frame.text[:12]))
            if is_pic:
                pics.append((l, t, w, h))
        for pr in pics:
            if pr[2] >= SW - 0.1:
                continue  # 全屏背景豁免
            if pr[1] < 1.40 and pr[3] < SH - 1.0:
                issues.append(f'页{idx} 图片进入标题区 top={pr[1]:.2f}')
            for tr, name in texts:
                ox = max(0, min(pr[0] + pr[2], tr[0] + tr[2]) - max(pr[0], tr[0]))
                oy = max(0, min(pr[1] + pr[3], tr[1] + tr[3]) - max(pr[1], tr[1]))
                if ox > 0.02 and oy > 0.02:
                    issues.append(f'页{idx} 图文重叠 {name} ({ox:.2f},{oy:.2f})')

    print('=' * 60)
    if issues:
        print(f'✗ 发现 {len(issues)} 个问题:')
        for i in issues:
            print('  ' + i)
        sys.exit(1)
    print('EXTREME_OK 7个极端场景全部通过（越界/贴边/重叠/标题区）')


if __name__ == '__main__':
    main()
