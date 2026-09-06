# -*- coding: utf-8 -*-
"""测试 富文本内容 + 插图 的导出链路（PPTX + HTML 预览 + 读回验证 + 插图兜底）"""
import sys
import tempfile
import os

sys.path.insert(0, '.')

from settings_module import Presentation, Slide
from core.pptx_builder import PPTXBuilder


def make_presentation():
    p = Presentation(title="软件工程生涯规划", theme="business", language="zh")
    slides = [
        Slide(
            page=1, title="个人定位",
            points=["兴趣方向：后端工程", "能力盘点：算法与系统", "价值观：长期主义"],
            detail="通过三年竞赛与项目历练，我明确了自身定位：以扎实的算法基础为底盘，聚焦后端与系统方向，兼顾工程落地能力，追求可持续的长期成长而非短期投机。",
            notes="开场点明个人画像，语气自信自然",
            layout="title_content",
            image_prompt="flat vector illustration of a young software engineer standing at a crossroads with three signposts",
        ),
        Slide(
            page=2, title="职业目标",
            points=["短期：夯实基础", "中期：技术骨干", "长期：架构师/技术专家"],
            detail="职业目标分三步走：短期 1-3 年内打牢工程基础并积累完整项目经验；中期 3-5 年成为团队核心技术骨干；长期向架构师或领域专家演进，形成技术壁垒。",
            notes="强调目标的递进逻辑",
            layout="title_content",
            image_prompt="three ascending steps roadmap with flag milestones, minimal flat style",
        ),
        Slide(
            page=3, title="发展路径",
            points=["技术路线：纵深", "管理路线：横向", "复合路线：T 型"],
            detail="发展路径存在三种选择：技术路线追求专业纵深，管理路线强调团队与资源整合，复合路线以 T 型能力兼顾两者。结合个人特质，优先 T 型发展。",
            notes="对比三种路径，给出个人选择",
            layout="two_content",
            image_prompt="T-shaped skills diagram comparing deep technical and broad management skills",
        ),
    ]
    for s in slides:
        p.add_slide(s)
    return p


p = make_presentation()
builder = PPTXBuilder()

tmpdir = tempfile.mkdtemp()
out = os.path.join(tmpdir, "test.pptx")

# 1. 生成兜底插图（无 API Key 时 Pillow 主题图），绑定到第 1 页
try:
    from core.image_generator import ImageGenerator
    img = ImageGenerator(client=None)._generate_fallback(
        "个人定位", tmpdir, "slide_1", "business"
    )
    if img and os.path.exists(img):
        p.slides[0].image_path = img
        print(f"✓ 兜底插图生成成功: {img} ({os.path.getsize(img)} 字节)")
    else:
        print("✗ 兜底插图生成失败")
except Exception as e:
    import traceback
    print(f"✗ 插图生成异常: {e}")
    traceback.print_exc()

# 2. 导出 PPTX
try:
    result = builder.build(p, out)
    size = os.path.getsize(out)
    print(f"✓ PPTX 导出成功: {result} ({size} 字节)")
except Exception as e:
    import traceback
    print(f"✗ PPTX 导出失败: {e}")
    traceback.print_exc()

# 3. HTML 预览（含 detail 与内嵌插图）
try:
    html = builder.preview_html(p)
    assert "软件工程生涯规划" in html
    assert "个人定位" in html
    assert "data:image/png;base64," in html, "预览未内嵌插图"
    assert "通过三年竞赛与项目历练" in html, "预览缺少 detail 正文"
    print(f"✓ HTML 预览生成成功 ({len(html)} 字符，含 detail + 内嵌插图)")
except Exception as e:
    print(f"✗ HTML 预览失败: {e}")

# 4. 预览窗口渲染
try:
    os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
    from PyQt6.QtWidgets import QApplication
    from ui.preview_widget import PreviewWidget
    app = QApplication([])
    w = PreviewWidget()
    w.update_preview(p)
    html_rendered = w.main_preview.toHtml()
    # 预览展示的是单页详情，不包含整个演示文稿标题
    assert "个人定位" in html_rendered, "缺少第一页标题"
    assert "data:image/png;base64," in html_rendered, "预览未内嵌插图"
    assert "通过三年竞赛与项目历练" in html_rendered, "预览缺少 detail 正文"
    print("✓ 预览窗口渲染成功（含正文 + 插图）")
except Exception as e:
    import traceback
    print(f"✗ 预览窗口失败: {e}")
    traceback.print_exc()

# 5. 读回 PPTX 验证（标题 + detail 正文 + 要点 + 图片均写入）
try:
    from pptx import Presentation as PptxRead
    pr = PptxRead(out)
    assert len(pr.slides) == 3, f"幻灯片数量错误: {len(pr.slides)}"

    extracted_text = []
    image_count = 0
    for s in pr.slides:
        for sh in s.shapes:
            if sh.shape_type == 13:  # PICTURE
                image_count += 1
            if sh.has_text_frame:
                t = sh.text_frame.text.strip()
                if t:
                    extracted_text.append(t)
    joined = "\n".join(extracted_text)

    for expect in ["个人定位", "职业目标", "发展路径",
                   "兴趣方向", "能力盘点", "价值观",
                   "短期", "中期", "长期",
                   "技术路线", "管理路线", "复合路线",
                   "通过三年竞赛与项目历练"]:
        assert expect in joined, f"缺失内容: {expect}"
    assert image_count >= 1, "PPTX 未插入插图图片"
    print(f"✓ PPTX 读回验证成功（3 页标题 + 9 个要点 + detail 正文 + {image_count} 张插图）")
except Exception as e:
    import traceback
    print(f"✗ PPTX 读回验证失败: {e}")
    traceback.print_exc()

print()
print("=== 富文本 + 插图导出链路测试完成 ===")
